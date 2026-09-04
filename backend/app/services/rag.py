from __future__ import annotations

import io
from pathlib import Path
from typing import Optional
from urllib.parse import urlparse

import httpx
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import Document, KnowledgeChunk
from app.services.embeddings import chunk_text, cosine_similarity, embed_texts


async def extract_text_from_bytes(filename: str, data: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if lower.endswith(".docx"):
        from docx import Document as DocxDocument

        doc = DocxDocument(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    # txt / csv / fallback
    return data.decode("utf-8", errors="ignore")


async def ingest_text(
    db: AsyncSession,
    agent_id: int,
    content: str,
    filename: str,
    source_type: str = "text",
    file_url: Optional[str] = None,
) -> Document:
    document = Document(
        agent_id=agent_id,
        filename=filename,
        source_type=source_type,
        file_url=file_url,
        status="processing",
    )
    db.add(document)
    await db.flush()

    chunks = chunk_text(content)
    embeddings = await embed_texts(chunks)
    for idx, (chunk, emb) in enumerate(zip(chunks, embeddings)):
        db.add(
            KnowledgeChunk(
                document_id=document.id,
                agent_id=agent_id,
                content=chunk,
                embedding=emb,
                meta={"chunk_index": idx, "source": filename},
            )
        )

    document.status = "ready" if chunks else "empty"
    await db.commit()
    await db.refresh(document)
    return document


async def ingest_upload(
    db: AsyncSession,
    agent_id: int,
    filename: str,
    data: bytes,
    upload_dir: Path,
) -> Document:
    upload_dir.mkdir(parents=True, exist_ok=True)
    dest = upload_dir / f"{agent_id}_{filename}"
    dest.write_bytes(data)
    text = await extract_text_from_bytes(filename, data)
    return await ingest_text(
        db,
        agent_id=agent_id,
        content=text,
        filename=filename,
        source_type="upload",
        file_url=str(dest),
    )


async def ingest_website(db: AsyncSession, agent_id: int, url: str) -> Document:
    async with httpx.AsyncClient(timeout=30.0, follow_redirects=True) as client:
        resp = await client.get(url)
        resp.raise_for_status()
        html = resp.text
    # Lightweight strip of tags
    import re

    text = re.sub(r"<script[\s\S]*?</script>", " ", html, flags=re.I)
    text = re.sub(r"<style[\s\S]*?</style>", " ", text, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    host = urlparse(url).netloc or "website"
    return await ingest_text(
        db,
        agent_id=agent_id,
        content=text[:50000],
        filename=host,
        source_type="website",
        file_url=url,
    )


async def search_knowledge(
    db: AsyncSession,
    agent_id: int,
    query: str,
    top_k: int = 4,
) -> list[dict]:
    result = await db.execute(
        select(KnowledgeChunk).where(KnowledgeChunk.agent_id == agent_id)
    )
    chunks = list(result.scalars().all())
    if not chunks:
        return []

    query_emb = (await embed_texts([query]))[0]
    scored: list[tuple[float, KnowledgeChunk]] = []
    for chunk in chunks:
        emb = chunk.embedding
        if emb is None:
            continue
        # JSON fallback may deserialize fine; Vector may be list-like
        emb_list = list(emb) if not isinstance(emb, list) else emb
        score = cosine_similarity(query_emb, emb_list)
        scored.append((score, chunk))

    scored.sort(key=lambda x: x[0], reverse=True)
    return [
        {
            "content": c.content,
            "score": round(score, 4),
            "metadata": c.meta or {},
        }
        for score, c in scored[:top_k]
        if score > 0.05
    ]
